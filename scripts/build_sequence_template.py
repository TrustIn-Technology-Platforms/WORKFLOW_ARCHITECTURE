"""Build the recruiter-facing sequence document template.

    python scripts/build_sequence_template.py
    python scripts/build_sequence_template.py path/to/out.docx

Two files come out: the annotated template, whose rules ride in Word comments
(a part of the file the reader never opens, so a recruiter who leaves them in
posts nothing extra), and a paste-only copy with one slot per section for a
content team handed finished copy.

The order is the one TrustIn's own documents already use (Sohaib's working
template, 2026-10-08): connection note, shared subject, three emails, InMail,
the Wellfound advert, then the two sections only the search reads. There is no
general advert because Wellfound is the only board the automation posts to;
the Wellfound section is the advert. Shape and reasons:
docs/templates/sequence-document.md.
"""

from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DEFAULT_OUT = Path(__file__).resolve().parent.parent / "docs" / "templates" / "sequence-document-template.docx"
PASTE_OUT = DEFAULT_OUT.with_name("sequence-document-paste-here.docx")

# Fixed so rebuilding an unchanged template produces an unchanged file.
_COMMENT_DATE = "2026-10-08T00:00:00Z"

TITLE_SHAPE = "[Role title] / [Stage + what the company is] / [On-site or Remote] / [Stack] / up to $[max]k + equity"


class _Comments:
    """Margin comments, written as raw OOXML: python-docx 1.1 cannot add them."""

    def __init__(self, document) -> None:
        self._part = XmlPart(
            PackURI("/word/comments.xml"),
            CT.WML_COMMENTS,
            parse_xml(f"<w:comments {nsdecls('w')}/>"),
            document.part.package,
        )
        document.part.relate_to(self._part, RT.COMMENTS)
        self._next_id = 0

    def add(self, paragraph, text: str) -> None:
        cid = str(self._next_id)
        self._next_id += 1

        comment = OxmlElement("w:comment")
        for name, value in (("w:id", cid), ("w:author", "TrustIn"),
                            ("w:initials", "TI"), ("w:date", _COMMENT_DATE)):
            comment.set(qn(name), value)
        for line in text.strip().split("\n"):
            p, r, t = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
            t.text = line
            t.set(qn("xml:space"), "preserve")
            r.append(t)
            p.append(r)
            comment.append(p)
        self._part.element.append(comment)

        p = paragraph._p
        start = OxmlElement("w:commentRangeStart")
        start.set(qn("w:id"), cid)
        end = OxmlElement("w:commentRangeEnd")
        end.set(qn("w:id"), cid)
        ref_run, ref = OxmlElement("w:r"), OxmlElement("w:commentReference")
        ref.set(qn("w:id"), cid)
        ref_run.append(ref)
        properties = p.find(qn("w:pPr"))
        if properties is not None:
            properties.addnext(start)
        else:
            p.insert(0, start)
        p.append(end)
        p.append(ref_run)


def build(path: Path) -> Path:
    document = docx.Document()
    comments = _Comments(document)

    def line(text: str, note: str = ""):
        paragraph = document.add_paragraph(text)
        if note:
            comments.add(paragraph, note)
        return paragraph

    def heading(text: str, note: str = ""):
        # A real Heading 2, not a bold line: bold is only read as a heading
        # above the first styled one, and a pasted Client JD brings styled
        # headings with it.
        paragraph = document.add_heading(text, level=2)
        if note:
            comments.add(paragraph, note)
        return paragraph

    def message(body: str):
        line("Hi {first_name},")
        line(body)
        line("Yours Sincerely,")

    heading(
        "LinkedIn Connection",
        """\
START HERE
1. Save this file as "Company - Role - Location.docx", e.g. "Acme - Founding Platform Engineer - SF.docx". That name becomes the campaign name on noon, Loxo and Juicebox, word for word.
2. Replace every [bracketed] line. Anything left in brackets is posted as written.
3. Never rename or restyle the blue section headings. To add one, copy an existing heading.

THIS SECTION: the note on the LinkedIn connection request (noon's first step). Keep it under 300 characters including the greeting; longer is cut off.
Merge fields everywhere: {first_name} and {company} (the candidate's current company). {ai_intro} on its own line is noon's AI opener; the other platforms remove that line.""",
    )
    line("Hi {first_name},")
    line("[Connection note, under 300 characters in total.]")

    heading(
        "Subject",
        """\
One subject line for the whole email thread and the InMail. Emails 2 and 3 reply under it, so they carry no subject of their own.
Keep the salary in it: "... / up to $275k + equity".""",
    )
    line(TITLE_SHAPE)

    heading(
        "Email1",
        """\
One message per section. Two greetings under one heading means two messages got merged, and the run reports it.
Emails 2 and 3 go out 3 days apart on Loxo. To change a gap, write it in the heading: "Email2 (wait 5 days)".
No rhetorical questions, and no lines about TrustIn or yourself.""",
    )
    line("Hi {first_name},")
    line("{ai_intro}")
    line("[Email 1 body.]")
    line("Yours Sincerely,")

    heading(
        "Email2",
        "A follow-up in the same thread. Mention {company} here if anywhere; it becomes the candidate's current employer on every platform.",
    )
    message("[Email 2 body.]")

    heading(
        "Email3",
        "Keep exactly three emails. noon has three email slots and fails the run if one is empty. A fourth email would reach Loxo and Juicebox only.",
    )
    message("[Email 3 body.]")

    heading(
        "Inmail",
        """\
The LinkedIn InMail (noon's InMail step). It uses the Subject above unless its first line is "Subject: ...".
Leave the section out and noon sends Email 1's text in its place.""",
    )
    message("[InMail body.]")

    heading(
        "Wellfound",
        """\
The job advert, anonymised. This is what Wellfound posts; nothing else in the document is an advert.
First line: the title. Put the salary in it ("up to $275k + equity") or Wellfound's salary boxes stay empty.
State experience as "N+ years" somewhere in the copy; Wellfound's experience box is filled from it.
Bold sub-headings such as "Candidate Background" and "Technical Skills" are fine and stay inside this section.
Location comes from the posting row, not from this text.""",
    )
    line(TITLE_SHAPE)
    line("[Opening paragraph: the company and the problem, anonymised.]")
    line("[What they need you to build.]")
    document.add_paragraph().add_run("Candidate Background").bold = True
    line("[Background the client wants, including \"N+ years\" where it applies.]")
    document.add_paragraph().add_run("Technical Skills").bold = True
    line("[The stack and the kind of systems.]")
    line("[Closing line: who they want to talk to.]")

    heading(
        "Recruiter Notes",
        """\
Internal. Only the AI that drafts the candidate search reads this; it is never posted and never reaches a candidate.
Write what you learned that the JD does not say: the real company behind a codename, what the hiring manager stressed, competitors to source from, deal-breakers.
Two lines are read as data when present:
  Company: <real employer>  - the company list is built around it
  Skills: Go, Kubernetes, Terraform  - fills the skills on Wellfound when the row has none
Everything else is free text. Delete the section if there is nothing to add.""",
    )
    line("Company: [Real company name, if the filename uses a codename]")
    line("Skills: [Comma-separated skills the client stressed]")
    line("[Anything learned on the call that the JD does not say.]")

    heading(
        "Client JD",
        """\
Always the last section, under exactly this heading ("JD" also works; "Jd or job description" does not, and the spec would be posted to Wellfound as part of the advert).
Paste the client's job description here exactly as they sent it, headings and all.
noon, Loxo and Juicebox build the candidate search from this, not from the advert. The advert is written to attract applicants and softens the years, the stack and the location; the JD states them.
Everything below this heading is read as the JD, so nothing else goes after it.""",
    )
    line("[Paste the client's job description here, unchanged.]")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def build_paste_version(path: Path) -> Path:
    """The same shape with one paste slot per section and no reading.

    For a content team handed finished copy: every slot says what to paste.
    No comments, because the rules live in the brief and nobody reads a margin
    while pasting.
    """
    document = docx.Document()

    def slot(text: str) -> None:
        run = document.add_paragraph().add_run(f"<< {text} >>")
        run.italic = True

    def bold(text: str) -> None:
        document.add_paragraph().add_run(text).bold = True

    for heading, what in (
        ("LinkedIn Connection", "paste the connection note, under 300 characters, greeting included"),
        ("Subject", "paste the one subject line: Role / Stage / On-site / Stack / up to $Nk + equity"),
        ("Email1", "paste email 1, greeting to sign-off; {ai_intro} on its own line where noon's opener goes"),
        ("Email2", "paste email 2"),
        ("Email3", "paste email 3"),
        ("Inmail", "paste the LinkedIn InMail"),
        ("Wellfound", None),
        ("Recruiter Notes", None),
        ("Client JD", "paste the client's job description exactly as received"),
    ):
        document.add_heading(heading, level=2)
        if heading == "Wellfound":
            slot("title line: Role / Stage / On-site / Stack / up to $Nk + equity")
            slot("paste the advert opening, anonymised")
            bold("Candidate Background")
            slot("paste the background wanted, with \"N+ years\" where it applies")
            bold("Technical Skills")
            slot("paste the stack and systems")
            continue
        if heading == "Recruiter Notes":
            document.add_paragraph("Company: << real company name, if the file uses a codename >>")
            document.add_paragraph("Skills: << comma-separated skills the client stressed >>")
            slot("anything learned on the call that the JD does not say; internal, never posted")
            continue
        slot(what)

    document.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    print(build(out))
    print(build_paste_version(PASTE_OUT if len(sys.argv) == 1 else out.with_name(out.stem + "-paste-here.docx")))
