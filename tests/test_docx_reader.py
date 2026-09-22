"""Reading a .docx into style-tagged blocks.

The case that matters here is the document shape `Client JD` created: the
recruiter's own sections are bold lines, and the client's spec is pasted in
underneath carrying its Word Heading styles. Both kinds of heading in one file.
"""

from __future__ import annotations

from io import BytesIO

import docx

from app.documents import parser
from app.documents.docx_reader import read_blocks


def _bold(document, text: str) -> None:
    document.add_paragraph().add_run(text).bold = True


def _real_document() -> bytes:
    """The 2026-09-22 document, reduced to the shape that broke it.

    Bold section names throughout, then a pasted client spec whose own headings
    carry real Heading styles - which is what every document looks like once
    `Client JD` is filled in.
    """
    d = docx.Document()

    _bold(d, "Email1")
    d.add_paragraph("Subject: Applied AI Engineer / LA / up to $320K")
    d.add_paragraph("Hi {{first_name}}, the first email body.")

    _bold(d, "Email2")
    d.add_paragraph("Hi {{first_name}}, the second email body.")

    _bold(d, "Email3")
    d.add_paragraph("Hi {{first_name}}, the third email body.")

    _bold(d, "InMail")
    d.add_paragraph("The InMail body.")

    _bold(d, "Wellfound")
    d.add_paragraph("Applied AI Engineer / Distributed Systems / LA")
    d.add_paragraph("Board copy for Wellfound.")

    _bold(d, "Original Job Description")
    # The pasted spec. These are the headings that used to suppress every bold
    # heading above them.
    d.add_heading("Applied AI Engineer", level=1)
    d.add_heading("About the Role", level=2)
    d.add_paragraph("Productionising LLM and agentic workloads.")
    d.add_heading("What We're Looking For", level=2)
    d.add_paragraph("8+ years, Kubernetes, model serving at scale.")

    buffer = BytesIO()
    d.save(buffer)
    return buffer.getvalue()


def test_a_pasted_spec_does_not_suppress_the_bold_headings_above_it():
    """A real document posted with zero email steps on 2026-09-22.

    `_promote_pseudo_headings` bailed out on `any(is_heading)`, and the only
    real headings in the file were the ones inside the pasted client spec. So
    none of the bold section names above it were promoted, the whole sequence
    collapsed into one 6663-char advert, and Loxo, Juicebox and noon each
    reported that the document produced no email steps.
    """
    blocks = read_blocks(_real_document())

    promoted = [b.text for b in blocks if b.is_heading]
    assert "Email1" in promoted, "the bold section names must survive the pasted spec"
    assert "Email2" in promoted and "Email3" in promoted

    # The spec's own headings are untouched - everything below the Client JD
    # heading is the client's text verbatim, headings and all.
    assert "About the Role" in promoted

    document = parser.parse_document(blocks)
    channels = [(e.channel, e.order) for e in document.emails]
    assert channels == [("email", 1), ("email", 2), ("email", 3), ("inmail", 1)]
    assert document.client_jd and "Kubernetes" in document.client_jd
    assert document.advert_for("wellfound").body_text


def test_a_heading_styled_document_is_left_alone():
    """The guard this replaces existed for a reason: where the author formats
    every section, a bold line inside prose is emphasis, not a section break."""
    d = docx.Document()
    d.add_heading("Email 1", level=2)
    d.add_paragraph("Subject: Hello")
    para = d.add_paragraph()
    para.add_run("Email 2").bold = True

    buffer = BytesIO()
    d.save(buffer)
    blocks = read_blocks(buffer.getvalue())

    promoted = [b.text for b in blocks if b.is_heading]
    assert promoted == ["Email 1"], "a bold line below a real heading stays body"
